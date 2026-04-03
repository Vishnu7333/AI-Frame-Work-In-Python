from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
import os
from datetime import datetime

# 🔹 Create Workbook
wb = Workbook()

# 🔹 Summary Sheet
summary_ws = wb.active
summary_ws.title = "Summary"

summary_ws.append([
    "Sl No",
    "Step",
    "Expected Result",
    "Actual Result",
    "Status",
    "Screenshot"
])

# Header style
for cell in summary_ws[1]:
    cell.font = Font(bold=True)

# 🔹 Screenshot Sheet
screenshot_ws = wb.create_sheet("Screenshots")

# 🔹 Counters
row_counter = 2
step_counter = 1


# ✅ MAIN LOG FUNCTION
def log_result(step, expected, actual, status, screenshot_path):
    global row_counter, step_counter

    screenshot_path = os.path.abspath(screenshot_path)

    print(f"📝 Step {step_counter}: {step}")
    print("📸 Screenshot:", screenshot_path)

    # 🔗 Hyperlink
    link_cell = f"#Screenshots!A{row_counter}"

    summary_ws.append([
        step_counter,
        step,
        expected,
        actual,
        status,
        f'=HYPERLINK("{link_cell}", "View")'
    ])

    # 🎨 Status color
    last_row = summary_ws.max_row
    status_cell = summary_ws[f"E{last_row}"]

    if status.upper() == "PASS":
        status_cell.fill = PatternFill(start_color="C6EFCE", fill_type="solid")
    else:
        status_cell.fill = PatternFill(start_color="FFC7CE", fill_type="solid")

    # 📸 Screenshot sheet
    screenshot_ws[f"A{row_counter}"] = f"{step_counter}. {step}"

    try:
        if os.path.exists(screenshot_path):
            img = ExcelImage(screenshot_path)
            img.width = 350
            img.height = 220
            screenshot_ws.add_image(img, f"B{row_counter}")
        else:
            screenshot_ws[f"B{row_counter}"] = "❌ Image not found"
            print("❌ File not found:", screenshot_path)

    except Exception as e:
        screenshot_ws[f"B{row_counter}"] = "❌ Error loading image"
        print("❌ Excel error:", e)

    row_counter += 15
    step_counter += 1


# ✅ SAVE REPORT
def save_report():
    try:
        os.makedirs("reports", exist_ok=True)

        # 🔥 UNIQUE FILE NAME (NEW EVERY RUN)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        file_path = os.path.abspath(f"reports/report_{timestamp}.xlsx")

        # 📊 TABLE FORMAT
        total_rows = summary_ws.max_row
        table_ref = f"A1:F{total_rows}"

        table = Table(displayName="TestReport", ref=table_ref)

        style = TableStyleInfo(
            name="TableStyleMedium9",
            showRowStripes=True,
            showColumnStripes=False
        )

        table.tableStyleInfo = style
        summary_ws.add_table(table)

        # 📏 AUTO COLUMN WIDTH
        for col in summary_ws.columns:
            max_length = 0
            col_letter = col[0].column_letter

            for cell in col:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass

            summary_ws.column_dimensions[col_letter].width = max_length + 2

        # 💾 SAVE
        wb.save(file_path)

        print("✅ Report saved:", file_path)

    except Exception as e:
        print("❌ Failed to save report:", e)